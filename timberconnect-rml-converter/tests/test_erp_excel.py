"""
Tests fuer den ERP-Excel-Parser (Herstellungsvorgang).

Baut die Arbeitsmappe in-memory mit openpyxl auf — kein Fixture-Binary noetig.
Ausfuehren:  python -m pytest tests/test_erp_excel.py
"""

import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from services.erp_excel_service import (  # noqa: E402
    ERPExcelError,
    parse_erp_excel,
    sniff_erp_workbook,
)
from services.file_detector import FileType, detect_file_type  # noqa: E402

PLATE_EPC = "urn:epc:id:sgtin:4012345.001001.25VA000001-110"
LAMELLA_1 = "urn:epc:id:sgtin:4098765.012345.LAM-0001"
LAMELLA_2 = "urn:epc:id:sgtin:4098765.012345.LAM-0002"
LOT_EPC = "urn:epc:class:lgtin:4098765.012345.CH-2025-031"


def _build_workbook(
    identity: str | None = PLATE_EPC,
    identity_input: str | None = f"{LAMELLA_1}; {LAMELLA_2}; {LOT_EPC}",
    with_ident_sheet: bool = True,
) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Vertriebsauftraege"
    ws["A1"] = "Vertriebsauftraege"
    ws["A2"], ws["B2"] = "Artikel", "BSP-200001 Pos. 1001 X-100/5S"
    ws["A3"], ws["B3"] = "Liefertermin", datetime(2025, 4, 1)
    ws["A4"], ws["B4"] = "Transportavis_erforderlich", "ja"

    ws = wb.create_sheet("Artikelkonfiguration")
    ws["A1"] = "Artikelkonfiguration"
    ws["A2"], ws["B2"] = "Holzart", "FI Fichte"
    ws["A3"], ws["B3"] = "Breite", 1325
    ws["A4"], ws["B4"] = "Nettovolumen_Produkt", 0.07
    ws["A5"], ws["B5"] = "Anzahl_der_Schichten_innerhalb_einer_BSP-Platte", 5

    ws = wb.create_sheet("Produkionsauftraege")
    ws["A1"] = "Produktionsauftraege"
    ws["A2"], ws["B2"] = "Ist-Menge", 1
    ws["A3"], ws["B3"] = "Ist-Beginndatum", "31.03.2025 00:00:00"

    if with_ident_sheet:
        ws = wb.create_sheet("Identifikation")
        ws["A1"] = "Identifikation"
        if identity is not None:
            ws["A2"], ws["B2"] = "Identity", identity
        if identity_input is not None:
            ws["A3"], ws["B3"] = "IdentityInput", identity_input

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_valid_workbook():
    content = _build_workbook()
    document, warnings = parse_erp_excel(content, "abc123", "test.xlsx")

    assert document["__docId"] == "abc123"
    ident = document["identification"]
    # Eine Datei = eine Platte. identity steht trotzdem als einelementiges
    # Array, weil der module-Treiber des EECC nur Array-Notation liest.
    assert ident["identity"] == [PLATE_EPC]
    assert ident["identityInput"] == [LAMELLA_1, LAMELLA_2, LOT_EPC]
    assert "manufacturings" not in ident

    assert document["salesOrder"]["artikel"] == "BSP-200001 Pos. 1001 X-100/5S"
    assert document["salesOrder"]["liefertermin"] == "2025-04-01"
    assert document["salesOrder"]["transportavis_erforderlich"] is True
    assert document["product"]["holzart_v3"] == "FI Fichte"
    assert document["product"]["anzahl_der_Schichten_innerhalb_einer_BSP_Platte"] == 5
    assert document["productionOrder"]["ist_Beginndatum"] == "2025-03-31"

    # Fehlende Blaetter (Lieferauftraege etc.) sind Warnungen, keine Fehler.
    assert any("Lieferauftraege" in w for w in warnings)


def test_missing_identity_fails():
    content = _build_workbook(identity=None)
    with pytest.raises(ERPExcelError, match="Identity"):
        parse_erp_excel(content, "abc123")


def test_missing_ident_sheet_fails():
    content = _build_workbook(with_ident_sheet=False)
    with pytest.raises(ERPExcelError, match="Identifikation"):
        parse_erp_excel(content, "abc123")


def test_invalid_epc_fails():
    content = _build_workbook(identity="keine-urn")
    with pytest.raises(ERPExcelError, match="Ungültige GS1-EPCs"):
        parse_erp_excel(content, "abc123")


def test_missing_input_is_warning_not_error():
    content = _build_workbook(identity_input=None)
    document, warnings = parse_erp_excel(content, "abc123")
    assert document["identification"]["identity"] == [PLATE_EPC]
    assert document["identification"]["identityInput"] == []
    assert any("IdentityInput" in w for w in warnings)


def test_multiple_identities_fail():
    # Eine ERP-Datei beschreibt genau EINE Platte — zwei Identity-EPCs
    # (egal ob in einer Zelle oder mehreren Zeilen) sind ein Fehler.
    second = PLATE_EPC.replace("-110", "-120")
    content = _build_workbook(identity=f"{PLATE_EPC}; {second}")
    with pytest.raises(ERPExcelError, match="genau"):
        parse_erp_excel(content, "abc123")


def test_sniff_and_detect():
    content = _build_workbook()
    matched, identity = sniff_erp_workbook(content)
    assert matched >= 3
    assert identity == [PLATE_EPC]

    detection = detect_file_type(content, "ERP_BSP_Eingabetabelle.xlsx")
    assert detection.file_type == FileType.ERP_XLSX
    assert detection.data_type == "herstellung"
    assert detection.mapping_id == "erp_bsp"
    assert detection.trace_id == PLATE_EPC
    assert detection.is_recognized


def test_detect_foreign_xlsx_unknown():
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Tabelle1"
    buf = BytesIO()
    wb.save(buf)

    detection = detect_file_type(buf.getvalue(), "irgendwas.xlsx")
    assert detection.file_type == FileType.UNKNOWN
    assert not detection.is_recognized

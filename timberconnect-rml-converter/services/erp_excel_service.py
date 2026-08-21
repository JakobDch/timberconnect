"""
ERP-Excel-Service (Herstellungsvorgang, BSP-Werk)

Liest die ERP-Exceldatei des Holzwerkstoffproduzenten (Vorlage
"ERP_BSP_Eingabetabelle.xlsx") und erzeugt daraus das maschinenlesbare
JSON-Zwischendokument, das vom RML-Mapping mappings/erp_bsp.rml.ttl in die
TimberConnect-Ontologie v6 materialisiert wird.

Aufbau der Datei: sieben fachliche Blaetter (Vertriebsauftraege,
Artikelkonfiguration, Produkionsauftraege [sic], Lieferauftraege, Kommissionen,
Transportauftraege, Ausgangsrechnungen abfragen), jeweils Schluessel/Wert-Paare
in Spalte A/B. Die Feldnamen entsprechen 1:1 den ERP-Properties der Ontologie
(skos:notation M-913 ff.).

Idente — Blatt "Identifikation":
    Identity       GS1-EPC des hergestellten Produkts (BSP-Platte). Vom ERP
                   aus der Auftragsnummer abgeleitet. PFLICHT — ohne
                   gueltigen Identity-EPC bricht die Konvertierung ab, denn
                   ohne Ident waere das Produkt im Datenraum nicht auffindbar.
    IdentityInput  GS1-EPCs des Vormaterials (Schnittholzlamellen/-lose aus
                   dem Aufsaegevorgang, SGTIN oder LGTIN). Mehrere Werte mit
                   ';' getrennt. Optional, aber ohne sie fehlt die
                   Verknuepfung zur vorherigen Wertschoepfungsstufe (Warnung).

    FESTGELEGT (08/2026): eine ERP-Datei beschreibt GENAU EINE BSP-Platte.
    Identity traegt deshalb genau einen EPC (mehrere sind ein Fehler), und
    die Datei entspricht als Ganzes genau einem EPCIS TransformationEvent:
    IdentityInput -> inputEPCList, Identity -> outputEPCList.
"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime, time
from io import BytesIO
from typing import Any, Optional

logger = logging.getLogger(__name__)


class ERPExcelError(Exception):
    """Fehler beim Lesen oder Validieren der ERP-Exceldatei."""


# Gleiche Validierung wie fuer die versteckten PDF-Formularfelder
# (pdf_template_service._EPC_URN_RE und viewer pdfIdentityService.ts).
_EPC_URN_RE = re.compile(
    r"^urn:epc:(id:sgtin|class:sgtin|class:lgtin):[0-9]+\.[0-9]+\.[A-Za-z0-9_-]+$"
)

# Trennzeichen fuer Mehrfachwerte — identisch zum PDF-Pfad.
_MULTI_VALUE_RE = re.compile(r"[;\s,]+")

IDENT_SHEET = "Identifikation"
IDENTITY_KEY = "Identity"
IDENTITY_INPUT_KEY = "IdentityInput"

# XLSX beginnt wie jedes ZIP mit "PK\x03\x04".
XLSX_MAGIC = b"PK\x03\x04"

# Werttypen: steuern die Normalisierung fuer das JSON (und damit die
# rr:datatype-Angaben im Mapping).
#   str      -> String
#   int      -> Ganzzahl
#   dec      -> Dezimalzahl
#   bool     -> "ja"/"nein" -> true/false
#   date     -> "yyyy-mm-dd"
#   datetime -> ISO 8601 (ERP liefert teils Strings "31.03.2025 00:00:00")
#   time     -> "hh:mm:ss"
#
# (Blattname, {Zeilenlabel: (jsonKey, Typ)}). Die jsonKeys entsprechen den
# lokalen Namen der Ontologie-Properties (tc:<jsonKey>), damit Mapping und
# Ontologie ohne Uebersetzungstabelle nebeneinander lesbar sind.
SHEET_FIELDS: dict[str, tuple[str, dict[str, tuple[str, str]]]] = {
    "Vertriebsauftraege": ("salesOrder", {
        "Artikel": ("artikel", "str"),
        "Bezug": ("bezug", "str"),
        "Kontraktposition": ("kontraktposition", "str"),
        "Liefertermin": ("liefertermin", "date"),
        "Versandtermin": ("versandtermin", "date"),
        "Wiedervorlagetermin": ("wiedervorlagetermin", "date"),
        "Transportavis_erforderlich": ("transportavis_erforderlich", "bool"),
        "gemaeß_statischer_Berechnung": ("gemaess_statischer_Berechnung", "bool"),
        "Pack-Gruppe": ("pack_Gruppe", "str"),
    }),
    "Artikelkonfiguration": ("product", {
        "Konfigurationsregel": ("konfigurationsregel", "str"),
        "Gesamtmenge": ("gesamtmenge", "int"),
        "Auftrag": ("auftrag", "str"),
        "Typ_der_BSP_Platte": ("typ_der_BSP_Platte", "str"),
        "Konstruktionsnummer": ("konstruktionsnummer", "str"),
        "Nettovolumen_Produkt": ("nettovolumen_Produkt", "dec"),
        "CNC-Abbund": ("cNC_Abbund", "bool"),
        "Zusatzabbund": ("zusatzabbund", "bool"),
        "Deckelage_BSP": ("deckelage_BSP", "str"),
        "Optik_Seite_1": ("optik_Seite_1", "str"),
        "Optik_Seite_2": ("optik_Seite_2", "str"),
        "Optik_Seite_3": ("optik_Seite_3", "str"),
        "Optik_Seite_4": ("optik_Seite_4", "str"),
        "Optik_Seite_5": ("optik_Seite_5", "str"),
        "Optik_Seite_6": ("optik_Seite_6", "str"),
        "Hoehe_/_Staerke": ("hoehe__Staerke", "dec"),
        "Breite": ("breite", "dec"),
        "Laenge": ("laenge_v3", "dec"),
        "Anzahl_der_Schichten_innerhalb_einer_BSP-Platte":
            ("anzahl_der_Schichten_innerhalb_einer_BSP_Platte", "int"),
        "Beschreibung_(Freitext)": ("beschreibung_Freitext", "str"),
        "Bauteil_Sonderform": ("bauteil_Sonderform", "str"),
        "Bruttovolumen_Produkt": ("bruttovolumen_Produkt", "dec"),
        "Produktnorm": ("produktnorm", "str"),
        "Festigkeit_/_Material_Produkt": ("festigkeit__Material_Produkt", "str"),
        "Brandschutzklasse_Produkt": ("brandschutzklasse_Produkt", "str"),
        "Holzart": ("holzart_v3", "str"),
        "Beschichtung_Seite_1_(BxL)": ("beschichtung_Seite_1_BxL", "str"),
        "Beschichtung_Seite_2_(HxL)": ("beschichtung_Seite_2_HxL", "str"),
        "Beschichtung_Seite_3_(BxH)": ("beschichtung_Seite_3_BxH", "str"),
        "Beschichtung_Seite_4_(BxH)": ("beschichtung_Seite_4_BxH", "str"),
        "Beschichtung_Seite_5_(HxL)": ("beschichtung_Seite_5_HxL", "str"),
        "Beschichtung_Seite_6_(BxL)": ("beschichtung_Seite_6_BxL", "str"),
        "Oberflaeche_einer_Plattenseite_": ("oberflaeche_einer_Plattenseite", "dec"),
        "Quadratmeter_aller_6_Plattenseiten":
            ("quadratmeter_aller_6_Plattenseiten", "dec"),
        "PA_Gruppe": ("pA_Gruppe", "str"),
        "Beschaffungsweg_(Eigenfertigung,_Zukauf)":
            ("beschaffungsweg_Eigenfertigung_Zukauf", "str"),
        "PEFC": ("pEFC", "bool"),
        "Merkmal,_ob_aus_Cadwork_importiert":
            ("merkmal_ob_aus_Cadwork_importiert", "bool"),
        "Eckausbildung": ("eckausbildung", "str"),
        "Produktionsstandort": ("produktionsstandort", "str"),
        "Vertriebsauftragsnummer": ("vertriebsauftragsnummer", "str"),
        "Organisation": ("organisation", "str"),
        "Verpackung": ("verpackung", "str"),
        "Kennzeichnung_der_Ware_(Label)": ("kennzeichnung_der_Ware_Label", "str"),
        "Beschichtung_Anzahl_Lagen_Seite_1": ("beschichtung_Anzahl_Lagen_Seite_1", "int"),
        "Beschichtung_Anzahl_Lagen_Seite_2": ("beschichtung_Anzahl_Lagen_Seite_2", "int"),
        "Beschichtung_Anzahl_Lagen_Seite_3": ("beschichtung_Anzahl_Lagen_Seite_3", "int"),
        "Beschichtung_Anzahl_Lagen_Seite_4": ("beschichtung_Anzahl_Lagen_Seite_4", "int"),
        "Beschichtung_Anzahl_Lagen_Seite_5": ("beschichtung_Anzahl_Lagen_Seite_5", "int"),
        "Beschichtung_Anzahl_Lagen_Seite_6": ("beschichtung_Anzahl_Lagen_Seite_6", "int"),
        "Bauteilstatus_Produktion": ("bauteilstatus_Produktion", "str"),
        "Plattencode_BSP": ("plattencode_BSP", "str"),
        "Meter": ("meter", "dec"),
        "Kubikmeter": ("kubikmeter", "str"),
        "CNC_Maschiene": ("cNC_Maschiene", "str"),
    }),
    "Produkionsauftraege": ("productionOrder", {
        "Produktionslagerort": ("produktionslagerort", "str"),
        "Lagerlogistikorganisation": ("lagerlogistikorganisation", "str"),
        "Vertriebs-Auftragsposition": ("vertriebs_Aufttagsposition", "str"),
        "Soll-Menge": ("soll_Menge", "int"),
        "Ist-Menge": ("ist_Menge", "int"),
        "Ausschussmenge": ("ausschussmenge", "int"),
        "Zugangslagerort": ("zugangslagerort", "str"),
        "Zustaendiger_Mitarbeiter": ("zustaendiger_Mitarbeiter", "str"),
        "Beladungsplan-Position": ("beladungsplan_Position", "str"),
        "Soll-Beginndatum": ("soll_Beginndatum", "date"),
        "Ist-Beginndatum": ("ist_Beginndatum", "date"),
        "Soll-Ruestzeit": ("soll_Ruestzeit", "dec"),
        "Ist-Ruestzeit": ("ist_Ruestzeit", "dec"),
    }),
    "Lieferauftraege": ("deliveryOrder", {
        "Lieferauftragsnummer": ("lieferauftragsnummer", "str"),
        "Lademittel": ("lademittel", "str"),
        "Nettogewicht": ("nettogewicht", "dec"),
        "Lagerort": ("lagerort", "str"),
        "Bestandseigentuemer": ("bestandseigentuemer", "str"),
        "Auspraegungstyp": ("auspraegungstyp", "str"),
        "Zu_liefernde_Menge_Stk": ("zu_liefernde_Menge_Stk", "int"),
        "Zu_liefernde_Menge_m³": ("zu_liefernde_Menge_m", "dec"),
        "Gelieferte/_rueckgemeldete_Menge_Stk.":
            ("gelieferte_rueckgemeldete_Menge_Stk", "int"),
        "Gelieferte/_rueckgemeldete_Menge_m³":
            ("gelieferte_rueckgemeldete_Menge_m", "dec"),
    }),
    "Kommissionen": ("commission", {
        "Lieferauftrag": ("lieferauftrag", "str"),
        "Lieferempfaenger": ("lieferempfaenger", "str"),
        "Spediteur": ("spediteur", "str"),
        "Lagerlogistikorgansiation": ("lagerlogistikorgansiation", "str"),
        "Lieferbedingung": ("lieferbedingung", "str"),
        "Versandbedingung": ("versandbedingung", "str"),
    }),
    "Transportauftraege": ("transportOrder", {
        "Lieferant": ("lieferant", "str"),
        "Route": ("route", "str"),
        "Startdatum": ("startdatum", "datetime"),
        "Lieferdatum": ("lieferdatum", "datetime"),
        "Status_Avisierung": ("status_Avisierung", "str"),
        "Erfasst_von": ("erfasst_von", "str"),
        "Erfasst_um": ("erfasst_um", "datetime"),
        "Belegdatum": ("belegdatum", "date"),
        "Transportmittel": ("transportmittel", "str"),
        "Entladeeinrichtung": ("entladeeinrichtung", "str"),
        "Info_VA": ("info_VA", "str"),
        "Zustaendiger_Mitarbeiter": ("zustaendiger_Mitarbeiter_v2", "str"),
        "Zuletzt_geaendert_von": ("zuletzt_geaendert_von", "str"),
        "Zuletzt_geaendert_um": ("zuletzt_geaendert_um", "date"),
        "Termin_fix": ("termin_fix", "str"),
        "Versandstelle": ("versandstelle", "str"),
        "Status_AV": ("status_AV", "str"),
        "Genehmigung_erforderlich": ("genehmigung_erforderlich", "bool"),
        "VA-Nr.": ("vA_Nr", "str"),
        "LT-Kennung": ("lT_Kennung", "str"),
        "QS-Kontrolle": ("qS_Kontrolle", "str"),
        "BV/_Kunde": ("bV_Kunde", "str"),
        "Hinweis_(intern)": ("hinweis_intern", "str"),
        "Abholung_Termin": ("abholung_Termin", "date"),
        "Abholung_Uhrzeit": ("abholung_Uhrzeit", "time"),
        "Abholung_PLZ": ("abholung_PLZ", "int"),
        "Abholung_Ort": ("abholung_Ort", "str"),
        "Lieferung_Termin": ("lieferung_Termin", "date"),
        "Lieferung_Uhrzeit": ("lieferung_Uhrzeit", "time"),
        "Lieferung_PLZ": ("lieferung_PLZ", "int"),
        "Lieferung_Ort": ("lieferung_Ort", "str"),
    }),
    "Ausgangsrechnungen abfragen": ("invoice", {
        "Rechnungsempfaenger": ("rechnungsempfaenger", "str"),
        "Rechnungsempfaenger_Name": ("rechnungsempfaenger_Name", "str"),
        "Rechnungsempfaenger_Adresse": ("rechnungsempfaenger_Adresse", "str"),
        "Belegdatum": ("belegdatum_v2", "date"),
        "Leistungsdatum": ("leistungsdatum", "date"),
        "Buchungsdatum": ("buchungsdatum", "date"),
        "Steuerschluessel": ("steuerschluessel", "dec"),
        "Artikelpreis-Klassifikation": ("artikelpreis_Klassifikation", "dec"),
    }),
}

# Fuer die Erkennung: je mehr dieser Blaetter vorhanden sind, desto sicherer
# ist es die ERP-BSP-Tabelle.
ERP_SHEET_NAMES = set(SHEET_FIELDS.keys()) | {IDENT_SHEET}


def _load_workbook(content: bytes):
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - Abhaengigkeit fehlt
        raise ERPExcelError(
            "openpyxl ist nicht installiert — Excel-Verarbeitung nicht verfügbar"
        ) from exc
    try:
        return openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ERPExcelError(f"Excel-Datei konnte nicht gelesen werden: {exc}") from exc


def sniff_erp_workbook(content: bytes) -> tuple[int, list[str]]:
    """Leichtgewichtige Erkennung fuer den file_detector.

    Returns:
        (Anzahl wiedererkannter Blattnamen, Identity-EPCs sofern vorhanden).
        (0, []) wenn die Datei kein lesbares XLSX ist.
    """
    if content[:4] != XLSX_MAGIC:
        return 0, []
    try:
        wb = _load_workbook(content)
    except ERPExcelError:
        return 0, []
    try:
        matched = sum(1 for name in wb.sheetnames if name in ERP_SHEET_NAMES)
        identity: list[str] = []
        if IDENT_SHEET in wb.sheetnames:
            pairs = _read_key_values(wb[IDENT_SHEET])
            for label, value in pairs:
                if label == IDENTITY_KEY:
                    identity.extend(_split_epcs(value))
        return matched, [e for e in identity if _EPC_URN_RE.match(e)]
    finally:
        wb.close()


def _read_key_values(ws) -> list[tuple[str, Any]]:
    """Spalte A = Label, Spalte B = Wert; leere Zeilen und Ueberschrift egal."""
    pairs: list[tuple[str, Any]] = []
    for row in ws.iter_rows(min_col=1, max_col=2):
        label = row[0].value if len(row) > 0 else None
        value = row[1].value if len(row) > 1 else None
        if label is None:
            continue
        pairs.append((str(label).strip(), value))
    return pairs


def _split_epcs(value: Any) -> list[str]:
    if value is None:
        return []
    return [part for part in _MULTI_VALUE_RE.split(str(value).strip()) if part]


def _normalize(value: Any, kind: str) -> Optional[Any]:
    """Zellwert -> JSON-Wert entsprechend rr:datatype im Mapping."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None

    if kind == "bool":
        text = str(value).strip().lower()
        if text in ("ja", "yes", "true", "1", "x"):
            return True
        if text in ("nein", "no", "false", "0"):
            return False
        return None
    if kind == "int":
        try:
            return int(float(str(value).replace(",", ".")))
        except (TypeError, ValueError):
            return None
    if kind == "dec":
        try:
            return float(str(value).replace(",", "."))
        except (TypeError, ValueError):
            return None
    if kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        # ERP-Strings wie "31.03.2025 00:00:00"
        for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(str(value), fmt).date().isoformat()
            except ValueError:
                continue
        return str(value)
    if kind == "datetime":
        if isinstance(value, datetime):
            return value.isoformat()
        for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(str(value), fmt).isoformat()
            except ValueError:
                continue
        return str(value)
    if kind == "time":
        if isinstance(value, time):
            return value.isoformat()
        if isinstance(value, datetime):
            return value.time().isoformat()
        return str(value)
    return str(value)


def _parse_identification(pairs: list[tuple[str, Any]]) -> tuple[str, list[str], list[str]]:
    """Identity/IdentityInput -> (Platten-EPC, Input-EPCs, Warnungen).

    Eine ERP-Datei beschreibt genau eine BSP-Platte: Identity muss genau
    einen gueltigen EPC enthalten. IdentityInput darf mehrere Werte tragen
    (';'-getrennt) und ist optional (Warnung statt Fehler).
    """
    outputs: list[str] = []
    inputs: list[str] = []
    invalid: list[str] = []

    for label, value in pairs:
        target = None
        if label == IDENTITY_KEY:
            target = outputs
        elif label == IDENTITY_INPUT_KEY:
            target = inputs
        if target is None:
            continue
        for epc in _split_epcs(value):
            if _EPC_URN_RE.match(epc):
                target.append(epc)
            else:
                invalid.append(f"{label}: '{epc}'")

    if invalid:
        raise ERPExcelError(
            "Ungültige GS1-EPCs im Blatt 'Identifikation' "
            "(erwartet urn:epc:id:sgtin:… / urn:epc:class:lgtin:…): "
            + "; ".join(invalid)
        )

    if not outputs:
        raise ERPExcelError(
            "Kein Produkt-Ident gefunden: Das Blatt 'Identifikation' muss im Feld "
            "'Identity' den GS1-EPC der hergestellten BSP-Platte enthalten "
            "(vom ERP aus der Auftragsnummer abgeleitet, Format "
            "urn:epc:id:sgtin:<gcp>.<itemref>.<serial>). Ohne diesen Ident wäre "
            "das Produkt im Datenraum nicht auffindbar."
        )
    if len(outputs) > 1:
        raise ERPExcelError(
            "Mehrere Identity-EPCs gefunden — eine ERP-Datei beschreibt genau "
            "EINE BSP-Platte. Für weitere Platten bitte je eine eigene Datei "
            "erzeugen. Gefunden: " + "; ".join(outputs)
        )

    warnings: list[str] = []
    if not inputs:
        warnings.append(
            "Blatt 'Identifikation': IdentityInput fehlt — die Verknüpfung zu "
            "den Lamellen-Identen des Aufsägevorgangs kann nicht hergestellt "
            "werden."
        )
    return outputs[0], inputs, warnings


def parse_erp_excel(
    content: bytes,
    doc_id: str,
    source_filename: str = "",
) -> tuple[dict, list[str]]:
    """ERP-Exceldatei -> JSON-Zwischendokument fuer mappings/erp_bsp.rml.ttl.

    Args:
        content: Rohbytes der .xlsx-Datei
        doc_id: kanonischer Dokument-Hash (SHA-256, 16 Hex) — Subjektbasis
        source_filename: Originaldateiname (nur Metadatum)

    Returns:
        (Dokument-Dict, nicht-blockierende Warnungen)

    Raises:
        ERPExcelError: Datei unlesbar, EPC ungueltig oder Identity fehlt.
    """
    wb = _load_workbook(content)
    try:
        if IDENT_SHEET not in wb.sheetnames:
            raise ERPExcelError(
                "Blatt 'Identifikation' fehlt: Die ERP-Exceldatei muss ein Blatt "
                "'Identifikation' mit den Feldern 'Identity' (EPC des Produkts) und "
                "'IdentityInput' (EPCs des Vormaterials aus dem Aufsägevorgang) "
                "enthalten — siehe docs/ERP_EXCEL.md."
            )

        identity, identity_input, warnings = _parse_identification(
            _read_key_values(wb[IDENT_SHEET])
        )

        # Die Datei als Ganzes entspricht genau einem EPCIS
        # TransformationEvent: identityInput -> inputEPCList,
        # identity -> outputEPCList.
        #
        # "identity" steht bewusst als Liste, obwohl eine ERP-Datei genau eine
        # BSP-Platte beschreibt (_parse_identification erzwingt das): Der
        # module-Treiber des EECC liest den Produkt-Ident ueber das Muster
        # "identification"."identity"-<n>, also ausschliesslich in
        # Array-Notation. Als Skalar geschrieben findet er ihn nicht und bricht
        # mit "affords EPCs or EPC class quantities" ab.
        document: dict[str, Any] = {
            "timberconnect_erp": {
                "format": "erp_bsp",
                "version": "1.0",
                "source_file": source_filename,
            },
            "__docId": doc_id,
            "identification": {
                "identity": [identity],
                "identityInput": identity_input,
            },
        }

        for sheet_name, (section, fields) in SHEET_FIELDS.items():
            if sheet_name not in wb.sheetnames:
                warnings.append(f"Blatt '{sheet_name}' fehlt — Abschnitt übersprungen.")
                continue
            data: dict[str, Any] = {}
            known = 0
            for label, value in _read_key_values(wb[sheet_name]):
                mapped = fields.get(label)
                if mapped is None:
                    continue
                known += 1
                key, kind = mapped
                normalized = _normalize(value, kind)
                if normalized is not None:
                    data[key] = normalized
            logger.info(
                "ERP-Blatt '%s': %d/%d bekannte Felder, %d mit Wert",
                sheet_name, known, len(fields), len(data),
            )
            if data:
                document[section] = data

        return document, warnings
    finally:
        wb.close()
